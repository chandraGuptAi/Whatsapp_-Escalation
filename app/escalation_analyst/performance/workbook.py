"""Native XLSX writer, original sheet layouts, formulas and reference style remapping."""

from __future__ import annotations
import calendar
from copy import deepcopy
from datetime import date, datetime, timezone
import io
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from ..logistics.workbook import safe
from .engine import SHEETS

REFERENCE = json.loads((Path(__file__).parent / "reference-format.json").read_text(encoding="utf-8"))
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", NS)


def export_workbook(
    result: dict, run: dict, selected: list[str] | None = None, columns: dict | None = None
) -> bytes:
    book = Workbook()
    book.remove(book.active)
    cache = {}
    formats = {}
    columns = columns or {}
    cutoff = date.fromisoformat(run["request"]["toDate"])
    for name in selected or SHEETS:
        data = result["sheets"][name]
        sheet = book.create_sheet(name)
        keep = columns.get(name) or list(range(len(data["columns"])))
        exact = name == "KAM_MTD_Performance" and name not in columns
        header = 7 if exact else 4
        if exact:
            for address, value in {
                "D2": "MONTH",
                "E2": cutoff.strftime("%b"),
                "G2": cutoff.year,
                "H2": f"SUMMARY {cutoff.strftime('%b').upper()}-{str(cutoff.year)[2:]}",
                "D3": "DAY PASSED",
                "E3": cutoff.day,
                "G3": cutoff,
                "D4": "DAYS TO GO",
                "E4": calendar.monthrange(cutoff.year, cutoff.month)[1] - cutoff.day,
                "B6": "KAM Emp MTD Performance",
                "G6": "Escalation",
                "K6": "Avg",
                "L6": "%",
                "M6": "Ranking",
            }.items():
                sheet[address] = value
            sheet.merge_cells("H2:M3")
            sheet.merge_cells("B6:F6")
            sheet.merge_cells("G6:J6")
        else:
            sheet.append([name.replace("_", " ")])
            sheet.append(
                [
                    f"{run['request']['fromDate']} to {run['request']['toDate']} | IST | Calculation {result['version']}"
                ]
            )
            sheet.append([data["note"]])
            for row in (1, 2, 3):
                if len(keep) > 1:
                    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(keep))
        formats[name] = {}
        for c, index in enumerate(keep, 1):
            sheet.cell(header, c, safe(data["columns"][index] or ""))
            sheet.column_dimensions[get_column_letter(c)].width = (
                24
                if c == 1
                else (
                    38
                    if data["columns"][index]
                    and any(
                        s in data["columns"][index].lower() for s in ("name", "remark", "definition", "value")
                    )
                    else 20
                )
            )
        for offset, values in enumerate(data["rows"], header + 1):
            for c, index in enumerate(keep, 1):
                value = values[index]
                value = ("", "Not available")[data["columns"][index] is not None] if value is None else value
                format = data["formats"].get(index, data["formats"].get(str(index)))
                if format == "date" and isinstance(value, str) and len(value) == 10 and value[4] == "-":
                    try:
                        value = date.fromisoformat(value)
                    except ValueError:
                        pass
                cell = sheet.cell(offset, c, safe(value))
                if format:
                    cell.number_format = "yyyy-mm-dd" if format == "date" else "0.0%"
                    formats[name][cell.coordinate] = cell.number_format
            if name not in columns:
                formulas = {}
                last = header + len(data["rows"])
                if name == "KAM_MTD_Performance":
                    formulas = {
                        11: f"IFERROR(I{offset}/H{offset}-1,0)",
                        12: f"RANK(I{offset},$I${header+1}:$I${last},0)",
                    }
                elif name == "Seller_MTD_Performance":
                    formulas = {
                        8: f"IFERROR(G{offset}/F{offset}-1,0)",
                        9: f"RANK(G{offset},$G${header+1}:$G${last},0)",
                    }
                elif name == "Backup":
                    formulas = {23: f"LEN(F{offset})", 19: f"S{offset}-R{offset}"}
                for index, formula in formulas.items():
                    if isinstance(values[index], (int, float)) and not isinstance(values[index], bool):
                        cell = sheet.cell(offset, index + 1, "=" + formula)
                        cache[(name, cell.coordinate)] = values[index]
        if exact:
            last = header + len(data["rows"])
            total = last + 1
            for index in range(1, len(data["columns"])):
                col = get_column_letter(index + 1)
                values = [row[index] for row in data["rows"]]
                if (
                    index >= 6
                    and index not in (10, 11, 12, 13, 26, 27, 28)
                    and values
                    and all(type(v) in (int, float) for v in values)
                ):
                    sheet.cell(total, index + 1, f"=SUM({col}8:{col}{last})")
                    cache[(name, f"{col}{total}")] = sum(values)
                    if 14 <= index <= 25:
                        sheet.cell(6, index + 1, f"=SUM({col}8:{col}{last})")
                        cache[(name, f"{col}6")] = sum(values)
            sheet.cell(total + 2, 2, data["note"])
            sheet.merge_cells(start_row=total + 2, start_column=2, end_row=total + 2, end_column=29)
        sheet.auto_filter.ref = f'{"B" if exact else "A"}{header}:{get_column_letter(len(keep))}{max(header,header+len(data["rows"]))}'
        sheet.freeze_panes = f"A{header+1}"
    book.properties.creator = "Bigship"
    book.properties.title = "Bigship Escalation Performance"
    content = io.BytesIO()
    book.save(content)
    return reference_styles(content.getvalue(), book.sheetnames, result, columns, cache, formats)


def reference_styles(
    content: bytes, names: list[str], result: dict, columns: dict, cache: dict, formats: dict
) -> bytes:
    styles = ET.fromstring(REFERENCE["styles"])
    xfs = styles.find(f"{{{NS}}}cellXfs")
    numfmts = styles.find(f"{{{NS}}}numFmts")
    if numfmts is None:
        numfmts = ET.Element(f"{{{NS}}}numFmts", {"count": "0"})
        styles.insert(0, numfmts)
    codes = {el.get("formatCode"): int(el.get("numFmtId")) for el in numfmts}
    next_format = max([163, *codes.values()]) + 1
    created = {}

    def formatted(base: int, code: str | None) -> int:
        nonlocal next_format
        if not code:
            return base
        key = (base, code)
        if key not in created:
            if code not in codes:
                codes[code] = next_format
                next_format += 1
                ET.SubElement(numfmts, f"{{{NS}}}numFmt", numFmtId=str(codes[code]), formatCode=code)
            el = deepcopy(xfs[base])
            el.set("numFmtId", str(codes[code]))
            created[key] = len(xfs)
            xfs.append(el)
        return created[key]

    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(
        output, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            raw = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                index = int(item.filename.split("sheet")[-1].split(".")[0]) - 1
                name = names[index]
                root = ET.fromstring(raw)
                profile = REFERENCE["sheets"].get(name, {})
                exact = name == "KAM_MTD_Performance" and name not in columns
                for cell in root.iter(f"{{{NS}}}c"):
                    address = cell.get("r")
                    col = "".join(c for c in address if c.isalpha())
                    row = int(address[len(col) :])
                    base = (
                        (profile.get("rows", {}).get(str(7 if row == 7 else 8), {}).get(col, 65))
                        if exact and row >= 7
                        else 80 if row == 1 else 72 if row == 4 else 63 if row < 4 else 65
                    )
                    if exact and row < 7:
                        base = profile.get("rows", {}).get(str(row), {}).get(col, 80)
                    cell.set("s", str(formatted(base, formats.get(name, {}).get(address))))
                    if (name, address) in cache:
                        value = cell.find(f"{{{NS}}}v")
                        if value is None:
                            value = ET.SubElement(cell, f"{{{NS}}}v")
                        value.text = str(cache[(name, address)])
                if exact and profile.get("cols"):
                    previous = root.find(f"{{{NS}}}cols")
                    if previous is not None:
                        root.remove(previous)
                    cols = ET.fromstring(profile["cols"])
                    for el in cols.iter():
                        el.tag = "{" + NS + "}" + el.tag.split("}")[-1]
                    data = root.find(f"{{{NS}}}sheetData")
                    root.insert(list(root).index(data), cols)
                raw = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            if item.filename == "xl/styles.xml":
                continue
            if item.filename == "xl/theme/theme1.xml":
                raw = REFERENCE["theme"].encode()
            target.writestr(item, raw)
        xfs.set("count", str(len(xfs)))
        numfmts.set("count", str(len(numfmts)))
        target.writestr("xl/styles.xml", ET.tostring(styles, encoding="utf-8", xml_declaration=True))
    return output.getvalue()
