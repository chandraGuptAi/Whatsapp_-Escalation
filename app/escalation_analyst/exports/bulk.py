"""Combine saved AI report tables into a single safe Excel workbook."""

from __future__ import annotations
import io
from openpyxl import Workbook
from .documents import tables
from ..logistics.workbook import safe, style_sheet
from ..errors import EscalationError


def export_bulk(reports: list[dict], runtime=None) -> bytes:
    if runtime is not None:
        return export_source_rows(reports, runtime)
    book = Workbook()
    sheet = book.active
    sheet.title = "Escalations"
    header = None
    source_header = None
    timing = ["Esc. Raised", "Revert TAT", "Revert Date", "Resolution Date"]
    for report in reports:
        found = tables(report["report"])
        if not found:
            continue
        table = found[0]
        if header is None:
            source_header = table[0]
            normalized_header = ["Esc. Raised" if value == "Escalation Raise Time" else value for value in source_header]
            header = ["Company", *normalized_header, *[name for name in timing if name not in normalized_header]]
            sheet.append(header)
        if table[0] != source_header:
            raise EscalationError(400, "Saved reports have incompatible table columns")
        for row in table[1:]:
            values = [report.get("group_name") or report["group_id"], *row]
            values.extend(["-"] * (len(header) - len(values)))
            sheet.append([v if v == "-" else safe(v) for v in values])
    if header is None:
        raise EscalationError(404, "No escalation tables available for export")
    style_sheet(sheet, 1, [30] * len(header))
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def export_source_rows(reports, runtime):
    from datetime import datetime
    from ..analysis.master_rows import HEADERS, report_rows
    from ..utils.ist import IST
    rows = report_rows(runtime, reports)
    if not rows:
        raise EscalationError(404, "No AWB enquiries available in retained messages for this report")
    book = Workbook()
    sheet = book.active
    sheet.title = "Escalations"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append([v if v == "-" else safe(v) for v in row["values"]])
        for column in (2,):
            cell = sheet.cell(sheet.max_row, column)
            if cell.value:
                cell.value = datetime.strptime(cell.value, "%d-%m-%Y %H:%M:%S")
                cell.number_format = "dd-mm-yyyy hh:mm:ss"
        for column, key, number_format in ((11, "raisedAt", "h:mm am/pm"),
                                            (13, "responseAt", "dd-mmm-yyyy"),
                                            (14, "resolutionAt", "dd-mmm-yyyy")):
            cell = sheet.cell(sheet.max_row, column)
            value = row[key]
            cell.value = datetime.fromtimestamp(value, IST).replace(tzinfo=None) if value is not None else "-"
            cell.number_format = number_format
        sheet.cell(sheet.max_row, 5).number_format = "@"
    style_sheet(sheet, 1, [36, 24, 48, 24, 24, 18, 50, 24, 16, 24, 24, 20, 24, 24])
    evidence = book.create_sheet("Evidence")
    evidence.append(["Group ID", "AWB", "Identity", "Seller message IDs", "Response message ID", "Resolution message ID", "Order status source"])
    for row in rows:
        evidence.append([safe(row["groupId"]), row["awb"], row["identityStatus"],
                         ", ".join(map(str, row["sourceMessageIds"])), row["responseMessageId"], row["resolutionMessageId"], row["orderStatusSource"]])
    style_sheet(evidence, 1, [36, 24, 20, 50, 24, 24, 32])
    notes = book.create_sheet("Report rules")
    notes.append(["Rule", "Meaning"])
    for rule in [("Time zone", "IST; dates are actual enquiry/reply timestamps."),
                 ("Count scope", "Saved report window; one seller message/AWB counts once. Staff replies do not increase the count."),
                 ("Identity", "ID and name must match master directory. Unknown/mismatched identities have blank seller counts."),
                 ("Replies", "Verified employee plus explicit AWB or quoted seller message. Ambiguous replies are not assigned."),
                 ("Order status", "ERP mapping as of report end, otherwise an explicit staff statement (unverified). Evidence sheet identifies the source."),
                 ("Resolution", "Mapped staff closure evidence or explicit receipt/resolution confirmation by the same verified seller after a linked staff reply. Bare thanks does not close an enquiry."),
                 ("Retention", "Only retained messages up to the saved report message cutoff are available.")]:
        notes.append(rule)
    style_sheet(notes, 1, [24, 110])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
