"""Media-to-text via Gemini, isolated from report generation; failures remain visible text."""

from __future__ import annotations

import base64
import csv
import io
import httpx
from openpyxl import load_workbook


def understand(media: dict | None, kind: str, settings) -> str:
    if not media:
        return f"[{kind}: analysis failed]"
    try:
        raw = base64.b64decode(media.get("data") or media.get("base64") or "", validate=True)
        # Keeps base64 requests below the existing AI Box 10MB request middleware.
        if len(raw) > 6 * 1024 * 1024:
            return f"[{kind}: too large to analyse ({len(raw)/1024/1024:.1f}MB)]"
        filename = media.get("filename") or ""
        mime = media.get("mimetype") or "application/octet-stream"
        if filename.lower().endswith((".csv", ".xlsx")):
            if filename.lower().endswith(".csv"):
                text = raw.decode("utf-8-sig", errors="replace")
            else:
                book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                lines = io.StringIO()
                writer = csv.writer(lines)
                for sheet in book:
                    writer.writerow([sheet.title])
                    for row in sheet.iter_rows(values_only=True):
                        writer.writerow(row)
                        if lines.tell() > 20000:
                            break
                book.close()
                text = lines.getvalue()
            return f"[spreadsheet {filename}]\n" + text[:20000]
        if not settings.gemini_key:
            return f"[{kind}: analysis failed; Gemini key not configured]"
        response = httpx.post(
            f"https://generativelanguage.googleapis.com/v1beta/models/{settings.gemini_model}:generateContent",
            headers={"x-goog-api-key": settings.gemini_key},
            json={
                "contents": [
                    {
                        "parts": [
                            {
                                "text": "Extract the operational text from this media. Do not invent information."
                            },
                            {
                                "inline_data": {
                                    "mime_type": "audio/ogg" if kind == "ptt" else mime,
                                    "data": base64.b64encode(raw).decode(),
                                }
                            },
                        ]
                    }
                ]
            },
            timeout=120,
        )
        response.raise_for_status()
        parts = response.json()["candidates"][0]["content"]["parts"]
        return (
            "[" + ("voice note" if kind == "ptt" else kind) + "] " + "".join(p.get("text", "") for p in parts)
        )
    except Exception:
        return f"[{kind}: analysis failed]"
