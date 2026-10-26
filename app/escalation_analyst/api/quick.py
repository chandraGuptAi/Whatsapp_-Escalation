"""Quick reports: a free-form prompt over selected groups, attached files and previous reports."""

from __future__ import annotations

import base64
import binascii
import io
import re
import uuid
import zipfile
from typing import Literal
from xml.etree import ElementTree

from fastapi import APIRouter, Depends, Query, Response
from pydantic import Field

from .dependencies import runtime, user
from .models import Input
from ..analysis.reports import prompt, transcript
from ..errors import EscalationError
from ..exports.bulk import export_bulk
from ..exports.documents import export
from ..utils.ist import resolve_window

router = APIRouter(dependencies=[Depends(user)], tags=["Escalation quick reports"])

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_FILE_CHARS = 60_000  # Per file, after extraction; keeps prompts inside the model context.
MAX_TRANSCRIPT_CHARS = 400_000
TEXT_TYPES = (".txt", ".csv", ".md", ".json", ".log", ".tsv")
SUPPORTED = TEXT_TYPES + (".xlsx", ".docx")


class QuickFile(Input):
    name: str = Field(min_length=1, max_length=200)
    data: str = Field(min_length=1, max_length=14_000_000)  # base64 of at most 10 MB


class QuickReport(Input):
    prompt: str = Field(min_length=3, max_length=4000)
    groupIds: list[str] = Field(default_factory=list, max_length=50)
    hours: float | None = Field(default=None, gt=0, le=8760)
    fromDate: str | None = None
    toDate: str | None = None
    contextReportIds: list[int] = Field(default_factory=list, max_length=5)
    outputFormat: Literal["web", "web_excel", "word", "detailed"] = "web"
    voice: bool = False
    files: list[QuickFile] = Field(default_factory=list, max_length=5)
    requestKey: str | None = Field(default=None, max_length=64)


def extract(file: QuickFile) -> str:
    name = file.name.lower()
    if not name.endswith(SUPPORTED):
        raise EscalationError(400, f"{file.name}: unsupported file type. Use TXT, CSV, Excel (.xlsx) or Word (.docx).")
    try:
        raw = base64.b64decode(file.data, validate=True)
    except (binascii.Error, ValueError):
        raise EscalationError(400, f"{file.name}: file data is not valid base64.")
    if len(raw) > MAX_FILE_BYTES:
        raise EscalationError(400, f"{file.name}: file is larger than 10 MB.")
    try:
        if name.endswith(TEXT_TYPES):
            text = raw.decode("utf-8-sig", errors="replace")
        elif name.endswith(".xlsx"):
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            lines = []
            for sheet in book.worksheets:
                lines.append(f"## Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    if any(cell is not None for cell in row):
                        lines.append(" | ".join("" if cell is None else str(cell) for cell in row))
            text = "\n".join(lines)
        else:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                root = ElementTree.fromstring(archive.read("word/document.xml"))
            ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            text = "\n".join(
                "".join(node.text or "" for node in para.iter(ns + "t")) for para in root.iter(ns + "p")
            )
    except EscalationError:
        raise
    except Exception:
        raise EscalationError(400, f"{file.name}: the file could not be read.")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise EscalationError(400, f"{file.name}: no readable text found.")
    return text[:MAX_FILE_CHARS] + ("\n...(file truncated)..." if len(text) > MAX_FILE_CHARS else "")


def generate(r, config: dict) -> dict | None:
    window = resolve_window(config)
    names = {}
    for group_id in config["groupIds"]:
        names[group_id] = r.groups.get(group_id)["group_name"]
    messages = r.groups.messages(None, {**window, "inclusive": True}, config["groupIds"]) if names else []
    if not messages and not config["files"]:
        return None
    sections = []
    for group_id, name in names.items():
        rows = [m for m in messages if m["group_id"] == group_id]
        sections.append(f"### Group: {name} ({len(rows)} messages)\n" + (transcript(rows) if rows else "(no messages in period)"))
    chats = "\n\n".join(sections)
    if len(chats) > MAX_TRANSCRIPT_CHARS:
        chats = "...(earlier messages trimmed)...\n" + chats[-MAX_TRANSCRIPT_CHARS:]
    user_text = f"Operator request:\n\"\"\"\n{config['prompt']}\n\"\"\"\n\nPeriod: {window['label']}\n"
    if config["outputFormat"] == "detailed":
        user_text += "Depth: detailed analysis. Cover every distinct issue with evidence.\n"
    for report_id in config["contextReportIds"]:
        previous = r.reports.get(report_id)
        user_text += f"\nPrevious report ({previous['window_label']}, {previous.get('group_name') or ''}):\n{previous['report'][:20000]}\n"
    for file in config["files"]:
        user_text += f"\nAttached file: {file['name']}\n\"\"\"\n{file['text']}\n\"\"\"\n"
    if names:
        user_text += f"\nGroup transcripts:\n\"\"\"\n{chats}\n\"\"\""
    system, _ = prompt("quick-report.md")
    text = r.provider.generate(system, user_text, dict(scope="quick", group_id="__quick__", group_name="Quick report"))
    return dict(report=text, messageCount=len(messages), windowLabel=window["label"])


def start(r, config: dict) -> dict:
    job, created = r.start("quick", "quick:" + config["runKey"], config, lambda _: generate(r, config))
    return dict(id=job["id"], reused=not created)


def public(job: dict, full: bool = False) -> dict:
    request, result = job["request"], job["result"] or {}
    row = dict(
        id=job["id"],
        status=job["status"],
        createdAt=job["created_at"],
        updatedAt=job["updated_at"],
        prompt=request["prompt"] if full else request["prompt"][:200],
        inputs=request["inputs"],
        groupIds=request["groupIds"],
        groupNames=request["groupNames"],
        fileNames=[f["name"] for f in request["files"]],
        outputFormat=request["outputFormat"],
        windowLabel=result.get("windowLabel"),
        messageCount=result.get("messageCount"),
        error=job["error"],
    )
    if full:
        row["report"] = result.get("report")
    return row


def owned(r, job_id: int, owner) -> dict:
    job = r.jobs.get(job_id, "quick")
    if job["request"].get("owner") != owner.id:
        raise EscalationError(404, "Run not found")
    return job


@router.post("/quick-reports", status_code=202)
def create(body: QuickReport, owner=Depends(user), r=Depends(runtime)) -> dict:
    if not body.groupIds and not body.files:
        raise EscalationError(400, "Select at least one group or attach a file.")
    if (body.fromDate or body.toDate) and body.hours:
        raise EscalationError(400, "Choose either a rolling period or custom dates.")
    groups = [r.groups.get(group_id) for group_id in dict.fromkeys(body.groupIds)]
    for report_id in body.contextReportIds:
        r.reports.get(report_id)
    files = [dict(name=f.name, text=extract(f)) for f in body.files]
    config = dict(
        owner=owner.id,
        prompt=body.prompt,
        groupIds=[g["group_id"] for g in groups],
        groupNames=[g["group_name"] for g in groups],
        hours=body.hours,
        fromDate=body.fromDate,
        toDate=body.toDate,
        contextReportIds=body.contextReportIds,
        outputFormat=body.outputFormat,
        files=files,
        inputs=[k for k, used in (("prompt", True), ("groups", groups), ("files", files), ("voice", body.voice)) if used],
        requestKey=body.requestKey,
        runKey=body.requestKey or uuid.uuid4().hex,
    )
    resolve_window(config)  # Reject invalid dates before a job is claimed.
    return start(r, config)


@router.get("/quick-reports")
def history(
    limit: int = Query(10, ge=1, le=50), offset: int = Query(0, ge=0), owner=Depends(user), r=Depends(runtime)
) -> dict:
    rows = [j for j in r.jobs.list("quick") if j["request"].get("owner") == owner.id]
    return dict(runs=[public(j) for j in rows[offset : offset + limit]], total=len(rows), limit=limit, offset=offset)


@router.get("/quick-reports/{job_id}")
def detail(job_id: int, owner=Depends(user), r=Depends(runtime)) -> dict:
    return {"run": public(owned(r, job_id, owner), full=True)}


@router.post("/quick-reports/{job_id}/retry", status_code=202)
def retry(job_id: int, owner=Depends(user), r=Depends(runtime)) -> dict:
    job = owned(r, job_id, owner)
    if job["status"] == "running":
        raise EscalationError(409, "This run is still processing.")
    return start(r, {**job["request"], "requestKey": None, "runKey": uuid.uuid4().hex})


@router.get("/quick-reports/{job_id}/download")
def download(
    job_id: int,
    format: Literal["doc", "md", "csv", "xlsx"] = "doc",
    owner=Depends(user),
    r=Depends(runtime),
) -> Response:
    text = (owned(r, job_id, owner)["result"] or {}).get("report")
    if not text:
        raise EscalationError(404, "This run has no report to download.")
    if format == "xlsx":
        content = export_bulk([dict(report=text, group_name="Quick report", group_id="quick")])
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        content, media = export(text, format)
    return Response(
        content,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="quick-report-{job_id}.{format}"'},
    )
