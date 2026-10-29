"""Media to text. Spreadsheets are read locally; images, documents and voice notes go to Gemini
(media only, never reports). Every Gemini call is recorded in app.ai_calls.
Ported from app/escalation_analyst/whatsapp/media.py of the old service."""
from __future__ import annotations

import csv
import hashlib
import io
import time
from uuid import UUID

from app.core.db import tenant_transaction

SPREADSHEET_LIMIT = 20000
INSTRUCTION = "Extract the operational text from this media. Do not invent information."


def spreadsheet_text(raw: bytes, filename: str) -> str | None:
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
    elif lower.endswith(".xlsx"):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        lines = io.StringIO()
        writer = csv.writer(lines)
        for sheet in book:
            writer.writerow([sheet.title])
            for row in sheet.iter_rows(values_only=True):
                writer.writerow(row)
                if lines.tell() > SPREADSHEET_LIMIT:
                    break
        book.close()
        text = lines.getvalue()
    else:
        return None
    return f"[spreadsheet {filename}]\n" + text[:SPREADSHEET_LIMIT]


async def extract_text(db, provider, tenant_id: UUID, raw: bytes, *, wa_kind: str, mime_type: str,
                       filename: str) -> tuple[str, str, str | None]:
    """Returns (text, extract_status, model)."""
    sheet = spreadsheet_text(raw, filename)
    if sheet is not None:
        return sheet, "done", None
    if not provider.api_key:
        return f"[{wa_kind}: analysis skipped; Gemini not configured]", "skipped", None
    mime = "audio/ogg" if wa_kind == "ptt" else mime_type
    started = time.monotonic()
    status, error, text, usage = "ok", None, "", {"tokens_in": 0, "tokens_out": 0}
    try:
        text, usage = await provider.media_to_text(raw, mime, INSTRUCTION)
    except Exception as exc:
        status, error = "error", f"{type(exc).__name__}: {exc}"[:500]
    async with tenant_transaction(db, tenant_id) as conn:
        await conn.execute(
            """INSERT INTO app.ai_calls (tenant_id, provider, model, purpose, prompt_version, payload_hash,
                                         fields_sent, tokens_in, tokens_out, latency_ms, status, error)
               VALUES (%s, 'gemini', %s, 'media_extract', 'media-v1', %s, %s, %s, %s, %s, %s, %s)""",
            (tenant_id, provider.model, hashlib.sha256(raw).hexdigest(), [mime], usage["tokens_in"],
             usage["tokens_out"], int((time.monotonic() - started) * 1000), status, error),
        )
    if status != "ok":
        return f"[{wa_kind}: analysis failed]", "failed", provider.model
    label = "voice note" if wa_kind == "ptt" else wa_kind
    return f"[{label}] {text}", "done", provider.model
