"""Evidence-aware logistics workbook, with safe cells and the original 16-column order."""

from __future__ import annotations

from datetime import date
import io
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from app.shared.ist import format_time
from app.shared.logistics.tat import format_duration

HEADERS = [
    "Date",
    "Group Name",
    "Processing Method",
    "Seller/Employee",
    "No. of Escalation Entries",
    "AWB Number",
    "Mile Stage",
    "Category / Remarks",
    "Current Status",
    "TAT",
    "Esc. Raised",
    "Esc. Closed",
    "Revert TAT",
    "Resolution TAT",
    "Revert Date",
    "Resolution Date",
]


def safe(value):
    return "'" + value if isinstance(value, str) and re.match(r"^\s*[=+@-]", value) else value


def style_sheet(sheet, header: int, widths: list[int]) -> None:
    for cell in sheet[header]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF087F6A")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = f"A{header+1}"
    sheet.auto_filter.ref = f"A{header}:{sheet.cell(max(header,sheet.max_row),len(widths)).coordinate}"
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width


def export_workbook(records: list[dict], request: dict) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Logistics Data"
    title = "UNIQUE AWB \u2014 ESCALATION PERFORMANCE & ACTION SHEET"
    book.properties.title = title
    book.properties.creator = "Bigship Escalation Analyst"
    sheet.append([title])
    sheet.append(
        [
            f"Period: {request['fromDate']} to {request['toDate']} | IST | Retained captured messages only; mentions are not verified incidents."
        ]
    )
    sheet.append(HEADERS)
    for r in records:
        rows = [
            date.fromisoformat(r["date"]),
            r["groupName"],
            "Legacy estimates" if r["mode"] == "legacy" else "Deterministic; evidence-aware",
            r.get("sender") or "Not available",
            r["entryCount"],
            r["awb"],
            r["mile"],
            r["category"],
            r["status"],
            format_duration(r["responseSeconds"]),
            format_time(r["firstAt"]),
            format_time(r["closureAt"]),
            (
                ("Legacy estimate: " + format_duration(r["legacyResponseSeconds"]))
                if r["mode"] == "legacy" and r["legacyResponseSeconds"] is not None
                else format_duration(r["responseSeconds"])
            ),
            format_duration(r["resolutionSeconds"]),
            format_time(r["legacyResponseAt"] if r["mode"] == "legacy" else r["responseAt"]),
            format_time(r["closureAt"]),
        ]
        sheet.append([safe(v) for v in rows])
        sheet.cell(sheet.max_row, 1).number_format = "yyyy-mm-dd"
    style_sheet(sheet, 3, [14, 36, 30, 30, 25, 22, 18, 44, 28, 20, 26, 26, 28, 22, 26, 26])
    evidence = book.create_sheet("Evidence")
    evidence.append(
        [
            "Group ID",
            "AWB",
            "Message ID",
            "Timestamp IST",
            "Sender",
            "Exact stored text",
            "Observed status",
            "Classification rule",
            "AWB confidence",
            "Legacy revert TAT (estimate)",
        ]
    )
    for r in records:
        for m, o in zip(r["sourceMessages"], r["observations"]):
            evidence.append(
                [
                    safe(v)
                    for v in [
                        r["groupId"],
                        r["awb"],
                        str(m["id"]),
                        format_time(m["timestamp"]),
                        m.get("author_name") or m.get("author") or "Not available",
                        m.get("body"),
                        o["status"],
                        o["evidence"] or "Unclassified",
                        o["extraction"]["confidence"],
                        format_duration(r["legacyResponseSeconds"]),
                    ]
                ]
            )
    style_sheet(evidence, 1, [30, 22, 16, 28, 30, 70, 25, 45, 20, 26])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
