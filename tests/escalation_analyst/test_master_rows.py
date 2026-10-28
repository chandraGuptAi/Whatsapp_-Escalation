import io
from datetime import datetime
from types import SimpleNamespace

from openpyxl import load_workbook

from app.escalation_analyst.analysis.master_rows import build_rows, identity, HEADERS
from app.escalation_analyst.logistics.awb import extract_awbs
from app.escalation_analyst.exports.bulk import export_bulk
from .test_core import client  # noqa: F401


MAPPINGS = dict(sellers=[dict(id="101", name="Acme"), dict(id="102", name="Other")],
                employees=[dict(id="7", name="Agent")], members=[], shipments=[])
REPORT = dict(group_id="g", group_name="Group", report="", period_start=100, period_end=500)


def msg(i, body, name="Acme-101-Seller", author="seller", **extra):
    return dict(id=i, wa_id=f"wa-{i}", group_id="g", timestamp=100 + i * 10,
                author=author, author_name=name, body=body, **extra)


def test_all_awbs_and_explicit_lists():
    text = r"AWB 9876543210 / 8765432109 \ BGSC0078363792, 13090326121011 | 321003495"
    refs = extract_awbs(text)
    assert [r["awb"] for r in refs] == ["9876543210", "8765432109", "BGSC0078363792", "13090326121011", "321003495"]
    assert all(r["confidence"] == "explicit" for r in refs)
    assert extract_awbs("Contact 9876543210") == []
    assert extract_awbs("12345678901234567") == []


def test_bare_phone_shaped_awb_is_retained_but_contact_number_is_not():
    rows = build_rows([msg(1, r"9876543210\8765432109 Contact: 7654321098")], REPORT, MAPPINGS)
    assert {row["awb"] for row in rows} == {"9876543210", "8765432109"}


def test_count_messages_not_occurrences_or_staff_or_duplicates():
    first = msg(1, "12345678 / 23456789 / 12345678")
    messages = [first, first, msg(2, r"12345678\BGSC0078363792"),
                msg(3, "AWB 12345678 in transit", "Agent-7-BIGSHIP", "staff")]
    rows = {r["awb"]: r for r in build_rows(messages, REPORT, MAPPINGS)}
    assert rows["12345678"]["values"][3] == 2
    assert rows["23456789"]["values"][3] == 1
    assert rows["BGSC0078363792"]["values"][3] == 1
    assert rows["12345678"]["responseMessageId"] == 3
    assert rows["23456789"]["responseMessageId"] is None


def test_quoted_reply_links_multiple_awbs():
    rows = build_rows([msg(1, "12345678 / 23456789"),
                       msg(2, "Checking with courier", "Agent-7-BIGSHIP", "staff", quoted_wa_id="wa-1")], REPORT, MAPPINGS)
    assert len(rows) == 2
    assert all(r["responseMessageId"] == 2 for r in rows)


def test_ambiguous_reply_and_separate_sellers():
    messages = [msg(1, "12345678"), msg(2, "12345678", "Other-102-Seller", "other"),
                msg(3, "12345678 in transit", "Agent-7-BIGSHIP", "staff")]
    rows = build_rows(messages, REPORT, MAPPINGS)
    assert len(rows) == 2
    assert all(r["responseMessageId"] is None for r in rows)
    messages.append(msg(4, "Checking", "Agent-7-BIGSHIP", "staff", quoted_wa_id="wa-1"))
    rows = build_rows(messages, REPORT, MAPPINGS)
    assert rows[0]["responseMessageId"] == 4
    assert rows[1]["responseMessageId"] is None


def test_identity_validation_never_trusts_suffix_alone():
    assert identity(msg(1, ""), MAPPINGS)["role"] == "seller"
    assert identity(msg(1, "", "Wrong-101-Seller"), MAPPINGS)["role"] == "mismatch"
    assert identity(msg(1, "", "Acme-101-BIGSHIP"), MAPPINGS)["role"] == "mismatch"
    assert identity(msg(1, "", "Agent-999-BIGSHIP"), MAPPINGS)["role"] == "unknown"
    rows = build_rows([msg(1, "12345678", "Wrong-101-Seller")], REPORT, MAPPINGS)
    assert rows[0]["values"][3] is None


def test_no_resolution_from_delivery_words():
    rows = build_rows([msg(1, "12345678"), msg(2, "12345678 delivered", "Agent-7-BIGSHIP", "staff")], REPORT, MAPPINGS)
    assert rows[0]["resolutionMessageId"] is None
    assert rows[0]["values"][-1] == "-"


def test_seller_confirmation_closes_without_inflating_enquiry_count():
    messages = [msg(1, "12345678"), msg(2, "12345678 delivered", "Agent-7-BIGSHIP", "staff"),
                msg(3, "Received the parcel", quoted_wa_id="wa-2")]
    row = build_rows(messages, REPORT, MAPPINGS)[0]
    assert row["values"][3] == 1
    assert row["resolutionMessageId"] == 3
    reopened = build_rows([*messages, msg(4, "12345678 still pending")], REPORT, MAPPINGS)[0]
    assert reopened["resolutionMessageId"] is None
    assert reopened["values"][3] == 2
    acknowledged = build_rows([messages[0], messages[1], msg(3, "Thanks!", quoted_wa_id="wa-2")], REPORT, MAPPINGS)[0]
    assert acknowledged["values"][3] == 1 and acknowledged["resolutionMessageId"] is None
    messages[-1] = msg(3, "Not received the parcel", quoted_wa_id="wa-2")
    assert build_rows(messages, REPORT, MAPPINGS)[0]["resolutionMessageId"] is None


def test_status_is_scoped_to_each_awb_and_requests_are_not_facts():
    messages = [msg(1, "12345678 / 23456789"),
                msg(2, "12345678 in transit; 23456789 out for delivery", "Agent-7-BIGSHIP", "staff")]
    rows = {r["awb"]: r for r in build_rows(messages, REPORT, MAPPINGS)}
    assert rows["12345678"]["values"][7] == "In-Transit"
    assert rows["23456789"]["values"][7] == "Out for Delivery"
    messages[-1] = msg(2, "12345678 will be delivered", "Agent-7-BIGSHIP", "staff")
    assert build_rows(messages, REPORT, MAPPINGS)[0]["values"][7] == "Unknown"


def test_export_uses_source_rows_and_real_dates():
    runtime = SimpleNamespace(db=SimpleNamespace(read=lambda *a: []),
                              groups=SimpleNamespace(messages=lambda *a: [msg(1, "BGSC0078363792")]))
    book = load_workbook(io.BytesIO(export_bulk([REPORT], runtime)))
    sheet = book["Escalations"]
    assert [c.value for c in sheet[1]] == HEADERS
    assert isinstance(sheet.cell(2, 2).value, datetime)
    assert sheet.cell(2, 5).value == "BGSC0078363792"
    assert sheet.cell(2, 4).value is None
    assert [sheet.cell(1, c).value for c in range(11, 15)] == ["Esc. Raised", "Revert TAT", "Revert Date", "Resolution Date"]
    assert sheet.cell(2, 11).number_format == "h:mm am/pm"
    assert sheet.cell(2, 13).value == "-" and sheet.cell(2, 14).value == "-"


def test_timing_columns_display_formats_and_legacy_export():
    row = build_rows([msg(1, "12345678"), msg(7, "12345678 in transit", "Agent-7-BIGSHIP", "staff")], REPORT, MAPPINGS)[0]
    assert row["values"][10].endswith("am")
    assert row["values"][11] == "1 Min"
    assert row["values"][12] == "01-Jan-1970"
    legacy = {**REPORT, "report": "| AWB | Status |\n| --- | --- |\n| 12345678 | open |"}
    sheet = load_workbook(io.BytesIO(export_bulk([legacy]))).active
    assert [c.value for c in sheet[1]][-4:] == ["Esc. Raised", "Revert TAT", "Revert Date", "Resolution Date"]
    assert [c.value for c in sheet[2]][-4:] == ["-", "-", "-", "-"]


def test_reply_date_remains_an_excel_date():
    import json
    messages = [msg(1, "12345678"), msg(7, "12345678 in transit", "Agent-7-BIGSHIP", "staff")]
    runtime = SimpleNamespace(db=SimpleNamespace(read=lambda *a: [{"payload": json.dumps(MAPPINGS)}]),
                              groups=SimpleNamespace(messages=lambda *a: messages))
    sheet = load_workbook(io.BytesIO(export_bulk([REPORT], runtime))).active
    assert sheet.cell(2, 12).value == "1 Min"
    assert isinstance(sheet.cell(2, 13).value, datetime)
    assert sheet.cell(2, 13).number_format == "dd-mmm-yyyy"


def test_directory_and_source_rows_api(client):
    assert client.put("/master/directory", json=MAPPINGS).status_code == 200
    assert client.get("/master/directory").json()["sellers"] == MAPPINGS["sellers"]
    runtime = client.app.state.runtime
    runtime.groups.ingest(dict(wa_id="wa-1", group_id="g", author="seller", author_name="Acme-101-Seller", body="12345678", timestamp=110))
    runtime.groups.ingest(dict(wa_id="wa-2", group_id="g", author="staff", author_name="Agent-7-BIGSHIP", body="Checking", timestamp=120, quoted_wa_id="wa-1"))
    child = runtime.reports.save({**REPORT, "scope": "group", "max_message_id": 2})
    master = runtime.reports.save({**REPORT, "group_id": "__master__:all", "scope": "master", "group_reports": {"g": child}})
    response = client.get(f"/master/reports/{master}/rows")
    assert response.status_code == 200
    row = response.json()["rows"][0]
    assert row["values"][3] == 1 and row["responseMessageId"] == 2
    assert client.get(f"/master/bulk-export?reportId={master}").status_code == 200
